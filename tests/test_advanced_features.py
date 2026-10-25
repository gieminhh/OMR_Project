"""Kiểm thử tính năng Chấm hàng loạt, Sinh mẫu phiếu A4/PDF và Xuất Excel 6 sheet."""

from pathlib import Path
import tempfile
import unittest
import openpyxl

from modules.advanced_export import export_session_report
from modules.batch_processor import BatchProcessor
from modules.db import close_connections
from modules.session_manager import SessionManager
from modules.sheet_generator import generate_sheet, export_sheet_to_pdf, export_sheet_to_image


class AdvancedFeaturesTests(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.tmp_dir.name)
        self.db_path = self.base_dir / "test.sqlite"
        self.mgr = SessionManager(self.db_path)

    def tearDown(self):
        close_connections(self.db_path)
        self.tmp_dir.cleanup()

    def test_sheet_generator_pdf(self):
        sheet_v1 = generate_sheet("V1", sbd="001122", exam_code="101")
        sheet_v2 = generate_sheet("V2", session_code="01", sbd="003344", exam_code="102")

        pdf_path = self.base_dir / "sheets.pdf"
        export_sheet_to_pdf([sheet_v1, sheet_v2], pdf_path)
        self.assertTrue(pdf_path.is_file())
        self.assertGreater(pdf_path.stat().st_size, 5000)

    def test_batch_processor_queue_and_management(self):
        self.mgr.create_session("01", "Kiểm tra", "Toán", "2026-10-25")
        proc = BatchProcessor(self.mgr, "01", self.base_dir / "out")

        img1 = export_sheet_to_image(generate_sheet("V2", session_code="01", sbd="000001", exam_code="101"), self.base_dir / "img1.png")
        img2 = export_sheet_to_image(generate_sheet("V2", session_code="01", sbd="000002", exam_code="101"), self.base_dir / "img2.png")

        # Thêm file
        added = proc.add_files([img1, img2])
        self.assertEqual(added, 2)
        self.assertEqual(len(proc.items), 2)

        # Loại bỏ 1 ảnh
        proc.remove_item(0)
        self.assertEqual(len(proc.items), 1)
        self.assertEqual(proc.items[0].image_path.name, "img2.png")

    def test_export_excel_6_sheets(self):
        self.mgr.create_session("01", "Kỳ thi thử", "Tiếng Anh", "2026-10-25", class_codes=["12A1", "12A2"])
        self.mgr.add_student("HS01", "000001", "Nguyễn Văn A", "12A1")
        self.mgr.add_student("HS02", "000002", "Trần Thị B", "12A2")

        # Set đáp án chuẩn
        self.mgr.set_answer_key("01", "101", ["A"] * 50)

        # Bài 1: Thành công điểm 10
        self.mgr.save_exam_result("01", "img1.png", ["A"] * 50, sbd="000001", exam_code="101")
        # Bài 2: Cần kiểm tra (thiếu mã đề)
        self.mgr.save_exam_result("01", "img2.png", ["B"] * 50, sbd="000002", exam_code="")

        excel_path = self.base_dir / "report.xlsx"
        export_session_report(self.mgr, "01", excel_path)
        self.assertTrue(excel_path.is_file())

        wb = openpyxl.load_workbook(excel_path)
        expected_sheets = [
            "Tổng hợp đợt thi",
            "Danh sách điểm",
            "Thống kê theo lớp",
            "Thống kê theo mã đề",
            "Phân loại kết quả",
            "Bài cần kiểm tra",
        ]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Kiểm tra sheet tổng hợp
        ws1 = wb["Tổng hợp đợt thi"]
        self.assertIn("KỲ THI THỬ", str(ws1["A1"].value))

        # Kiểm tra sheet bài cần kiểm tra
        ws6 = wb["Bài cần kiểm tra"]
        self.assertEqual(ws6.cell(row=2, column=5).value, "NEEDS_REVIEW")


if __name__ == "__main__":
    unittest.main()
