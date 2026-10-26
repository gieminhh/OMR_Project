"""Mở rộng xuất báo cáo thống kê kỳ thi ra file Excel đa sheet (6 sheet) — Phase 10.

Các sheet:
1. Tổng hợp đợt thi
2. Danh sách điểm
3. Thống kê theo lớp
4. Thống kê theo mã đề
5. Phân loại kết quả
6. Danh sách bài cần kiểm tra
"""

from collections import Counter
from pathlib import Path
from typing import Any

from .errors import OMRError
from .session_manager import SessionManager, DEFAULT_GRADE_SCALE, classify_score

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False

# Màu sắc giao diện Excel
COLOR_HEADER_BG = "2E75B6"      # Xanh đậm
COLOR_HEADER_FG = "FFFFFF"      # Chữ trắng
COLOR_SECTION_BG = "D9E1F2"     # Xanh nhạt
COLOR_ALERT_BG = "FFF2CC"       # Vàng nhạt cảnh báo
COLOR_DANGER_BG = "FFC7CE"      # Đỏ nhạt lỗi
COLOR_SUCCESS_BG = "C6EFCE"     # Xanh lá


def _border() -> Border:
    s = Side(style="thin", color="D3D3D3")
    return Border(left=s, right=s, top=s, bottom=s)


def _apply_header_style(cell, text: str):
    cell.value = text
    cell.font = Font(bold=True, color=COLOR_HEADER_FG, name="Segoe UI", size=11)
    cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _border()


def _auto_fit_columns(ws, max_cols=15):
    for col in range(1, max_cols + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for cell in ws[col_letter]:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def export_session_report(
    session_manager: SessionManager,
    session_code: str,
    output_path: str | Path,
) -> Path:
    """Xuất toàn bộ báo cáo đợt thi ra file Excel 6 sheet."""
    if not _HAS_OPENPYXL:
        raise OMRError("Thư viện openpyxl chưa được cài đặt.")

    sess = session_manager.get_session(session_code)
    if not sess:
        raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")

    results = session_manager.list_results(session_code)
    out_path = Path(output_path)
    if out_path.suffix.lower() != ".xlsx":
        out_path = out_path.with_suffix(".xlsx")
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    # Xóa sheet mặc định
    wb.remove(wb.active)

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    bold = Font(bold=True, name="Segoe UI", size=11)
    thin_border = _border()

    # Thống kê cơ bản
    total_sheets = len(results)
    success_items = [r for r in results if r["status"] == "SUCCESS"]
    review_items = [r for r in results if r["status"] == "NEEDS_REVIEW"]
    failed_items = [r for r in results if r["status"] == "FAILED"]

    valid_scores = [r["score"] for r in success_items if r["score"] is not None]
    max_score = max(valid_scores) if valid_scores else 0.0
    min_score = min(valid_scores) if valid_scores else 0.0
    avg_score = round(sum(valid_scores) / len(valid_scores), 2) if valid_scores else 0.0

    # ──────────────────────────────────────────────────────────────────────────
    # SHEET 1: Tổng hợp đợt thi
    # ──────────────────────────────────────────────────────────────────────────
    ws1 = wb.create_sheet(title="Tổng hợp đợt thi")
    ws1.views.sheetView[0].showGridLines = True

    # Tiêu đề lớn
    ws1.merge_cells("A1:D1")
    title_cell = ws1["A1"]
    title_cell.value = f"BÁO CÁO TỔNG KẾT ĐỢT THI: {sess['name'].upper()}"
    title_cell.font = Font(bold=True, size=16, color=COLOR_HEADER_BG, name="Segoe UI")
    title_cell.alignment = center
    ws1.row_dimensions[1].height = 36

    info_data = [
        ("Mã đợt thi", sess["session_code"]),
        ("Tên đợt thi", sess["name"]),
        ("Môn thi", sess["subject"]),
        ("Ngày thi", sess["exam_date"]),
        ("Số lượng câu hỏi", sess["question_count"]),
        ("Thang điểm", sess["max_score"]),
        ("Tổng số bài đã chấm", total_sheets),
        ("Số bài hợp lệ (Chính thức)", len(success_items)),
        ("Số bài cần kiểm tra thủ công", len(review_items)),
        ("Số bài lỗi xử lý (Thất bại)", len(failed_items)),
        ("Điểm trung bình (bài hợp lệ)", avg_score),
        ("Điểm cao nhất", max_score),
        ("Điểm thấp nhất", min_score),
    ]

    for r_idx, (lbl, val) in enumerate(info_data, start=3):
        c1 = ws1.cell(row=r_idx, column=2, value=lbl)
        c2 = ws1.cell(row=r_idx, column=3, value=val)
        c1.font = bold
        c1.fill = PatternFill("solid", fgColor=COLOR_SECTION_BG)
        c1.border = thin_border
        c2.font = Font(name="Segoe UI", size=11)
        c2.alignment = center
        c2.border = thin_border
        if lbl in ("Điểm trung bình (bài hợp lệ)", "Điểm cao nhất"):
            c2.font = Font(bold=True, size=12, color="C00000", name="Segoe UI")

    _auto_fit_columns(ws1, 4)

    # ──────────────────────────────────────────────────────────────────────────
    # SHEET 2: Danh sách điểm
    # ──────────────────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet(title="Danh sách điểm")
    ws2.views.sheetView[0].showGridLines = True

    h2 = ["STT", "Số báo danh", "Họ và tên", "Lớp", "Đợt thi", "Mã đề",
          "Số câu đúng", "Số câu sai", "Bỏ trống", "Tô nhiều", "Điểm (thang 10)", "Xếp loại", "Trạng thái"]
    for c_idx, text in enumerate(h2, start=1):
        _apply_header_style(ws2.cell(row=1, column=c_idx), text)
    ws2.row_dimensions[1].height = 24

    for r_idx, item in enumerate(results, start=2):
        row_vals = [
            r_idx - 1,
            item.get("sbd") or "—",
            item.get("student_name") or "—",
            item.get("class_code") or "—",
            item["session_code"],
            item.get("exam_code") or "—",
            item.get("correct_count", "—"),
            item.get("wrong_count", "—"),
            item.get("blank_count", "—"),
            item.get("multiple_count", "—"),
            item.get("score") if item.get("score") is not None else "—",
            item.get("classification") or "—",
            item["status"],
        ]
        status_color = COLOR_SUCCESS_BG if item["status"] == "SUCCESS" else COLOR_ALERT_BG if item["status"] == "NEEDS_REVIEW" else COLOR_DANGER_BG
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = left if c_idx == 3 else center
            cell.border = thin_border
            # Giữ nguyên số 0 ở đầu cho SBD (cột 2), Đợt thi (cột 5), Mã đề (cột 6)
            if c_idx in (2, 5, 6) and val != "—":
                cell.number_format = '@'
                cell.value = str(val)
            if c_idx == 11 and isinstance(val, (int, float)):
                cell.font = Font(bold=True, color="C00000", name="Segoe UI")
            if c_idx == 13:
                cell.fill = PatternFill("solid", fgColor=status_color)
                cell.font = bold

    _auto_fit_columns(ws2, len(h2))

    # ──────────────────────────────────────────────────────────────────────────
    # SHEET 3: Thống kê theo lớp
    # ──────────────────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet(title="Thống kê theo lớp")
    ws3.views.sheetView[0].showGridLines = True

    h3 = ["Mã lớp", "Tổng số bài", "Bài hợp lệ", "Điểm TB", "Điểm cao nhất", "Điểm thấp nhất",
          "Giỏi (%)", "Khá (%)", "Trung bình (%)", "Yếu (%)", "Kém (%)"]
    for c_idx, text in enumerate(h3, start=1):
        _apply_header_style(ws3.cell(row=1, column=c_idx), text)
    ws3.row_dimensions[1].height = 24

    # Nhóm theo lớp (lấy cả bài SUCCESS)
    classes = sorted(list(set(r.get("class_code") or "Chưa phân lớp" for r in results)))
    for r_idx, ccode in enumerate(classes, start=2):
        class_items = [r for r in results if (r.get("class_code") or "Chưa phân lớp") == ccode]
        c_valid = [r for r in class_items if r["status"] == "SUCCESS" and r.get("score") is not None]
        c_scores = [r["score"] for r in c_valid]

        c_avg = round(sum(c_scores) / len(c_scores), 2) if c_scores else 0.0
        c_max = max(c_scores) if c_scores else 0.0
        c_min = min(c_scores) if c_scores else 0.0

        n_valid = len(c_valid)
        counts = Counter(r.get("classification") for r in c_valid)
        g_pct = f"{counts['Giỏi']/n_valid*100:.1f}%" if n_valid else "0%"
        k_pct = f"{counts['Khá']/n_valid*100:.1f}%" if n_valid else "0%"
        tb_pct = f"{counts['Trung bình']/n_valid*100:.1f}%" if n_valid else "0%"
        y_pct = f"{counts['Yếu']/n_valid*100:.1f}%" if n_valid else "0%"
        kem_pct = f"{counts['Kém']/n_valid*100:.1f}%" if n_valid else "0%"

        vals = [ccode, len(class_items), n_valid, c_avg, c_max, c_min, g_pct, k_pct, tb_pct, y_pct, kem_pct]
        for col_idx, v in enumerate(vals, start=1):
            cell = ws3.cell(row=r_idx, column=col_idx, value=v)
            cell.alignment = center
            cell.border = thin_border
            if col_idx == 4:
                cell.font = bold

    _auto_fit_columns(ws3, len(h3))

    # ──────────────────────────────────────────────────────────────────────────
    # SHEET 4: Thống kê theo mã đề
    # ──────────────────────────────────────────────────────────────────────────
    ws4 = wb.create_sheet(title="Thống kê theo mã đề")
    ws4.views.sheetView[0].showGridLines = True

    h4 = ["Mã đề", "Số bài làm", "Số bài hợp lệ", "Điểm TB", "Điểm cao nhất", "Điểm thấp nhất", "Số câu đúng TB"]
    for c_idx, text in enumerate(h4, start=1):
        _apply_header_style(ws4.cell(row=1, column=c_idx), text)
    ws4.row_dimensions[1].height = 24

    exam_codes = sorted(list(set(r.get("exam_code") or "Không rõ" for r in results)))
    for r_idx, ecode in enumerate(exam_codes, start=2):
        code_items = [r for r in results if (r.get("exam_code") or "Không rõ") == ecode]
        c_valid = [r for r in code_items if r["status"] == "SUCCESS" and r.get("score") is not None]
        c_scores = [r["score"] for r in c_valid]
        c_corrects = [r["correct_count"] for r in c_valid if r.get("correct_count") is not None]

        c_avg = round(sum(c_scores) / len(c_scores), 2) if c_scores else 0.0
        c_max = max(c_scores) if c_scores else 0.0
        c_min = min(c_scores) if c_scores else 0.0
        c_cor_avg = round(sum(c_corrects) / len(c_corrects), 1) if c_corrects else 0.0

        vals = [ecode, len(code_items), len(c_valid), c_avg, c_max, c_min, c_cor_avg]
        for col_idx, v in enumerate(vals, start=1):
            cell = ws4.cell(row=r_idx, column=col_idx, value=v)
            cell.alignment = center
            cell.border = thin_border
            if col_idx == 4:
                cell.font = bold

    _auto_fit_columns(ws4, len(h4))

    # ──────────────────────────────────────────────────────────────────────────
    # SHEET 5: Phân loại kết quả
    # ──────────────────────────────────────────────────────────────────────────
    ws5 = wb.create_sheet(title="Phân loại kết quả")
    ws5.views.sheetView[0].showGridLines = True

    h5 = ["Xếp loại", "Tiêu chuẩn thang 10", "Số lượng học sinh", "Tỷ lệ (%)"]
    for c_idx, text in enumerate(h5, start=1):
        _apply_header_style(ws5.cell(row=1, column=c_idx), text)
    ws5.row_dimensions[1].height = 24

    cfg = sess.get("grade_scale_config") or DEFAULT_GRADE_SCALE
    n_valid = len(valid_scores)
    counts = Counter(r.get("classification") for r in success_items)

    class_info = [
        ("Giỏi", f"Từ {cfg.get('gioi', 8.0):.1f} điểm", counts.get("Giỏi", 0)),
        ("Khá", f"Từ {cfg.get('kha', 6.5):.1f} đến dưới {cfg.get('gioi', 8.0):.1f}", counts.get("Khá", 0)),
        ("Trung bình", f"Từ {cfg.get('trung_binh', 5.0):.1f} đến dưới {cfg.get('kha', 6.5):.1f}", counts.get("Trung bình", 0)),
        ("Yếu", f"Từ {cfg.get('yeu', 3.5):.1f} đến dưới {cfg.get('trung_binh', 5.0):.1f}", counts.get("Yếu", 0)),
        ("Kém", f"Dưới {cfg.get('yeu', 3.5):.1f} điểm", counts.get("Kém", 0)),
    ]

    for r_idx, (name, criteria, cnt) in enumerate(class_info, start=2):
        pct = f"{cnt/n_valid*100:.1f}%" if n_valid else "0%"
        vals = [name, criteria, cnt, pct]
        for col_idx, v in enumerate(vals, start=1):
            cell = ws5.cell(row=r_idx, column=col_idx, value=v)
            cell.alignment = center
            cell.border = thin_border
            if col_idx in (1, 3):
                cell.font = bold

    _auto_fit_columns(ws5, len(h5))

    # ──────────────────────────────────────────────────────────────────────────
    # SHEET 6: Danh sách bài cần kiểm tra
    # ──────────────────────────────────────────────────────────────────────────
    ws6 = wb.create_sheet(title="Bài cần kiểm tra")
    ws6.views.sheetView[0].showGridLines = True

    h6 = ["STT", "Ảnh bài thi", "Số báo danh", "Mã đề", "Trạng thái", "Lý do cần kiểm tra"]
    for c_idx, text in enumerate(h6, start=1):
        _apply_header_style(ws6.cell(row=1, column=c_idx), text)
    ws6.row_dimensions[1].height = 24

    flagged_items = [r for r in results if r["status"] != "SUCCESS"]
    if not flagged_items:
        ws6.merge_cells("A2:F2")
        c = ws6["A2"]
        c.value = "Tất cả các bài thi đều hợp lệ. Không có bài nào cần kiểm tra!"
        c.font = Font(italic=True, color="008000", name="Segoe UI")
        c.alignment = center
    else:
        for r_idx, item in enumerate(flagged_items, start=2):
            fname = Path(item["image_path"]).name
            row_vals = [
                r_idx - 1,
                fname,
                item.get("sbd") or "Thiếu/Không đọc được",
                item.get("exam_code") or "Thiếu/Không đọc được",
                item["status"],
                item.get("review_reason") or "Lỗi nhận diện",
            ]
            fill_c = COLOR_ALERT_BG if item["status"] == "NEEDS_REVIEW" else COLOR_DANGER_BG
            for c_idx, v in enumerate(row_vals, start=1):
                cell = ws6.cell(row=r_idx, column=c_idx, value=v)
                cell.alignment = left if c_idx in (2, 6) else center
                cell.border = thin_border
                if c_idx in (3, 4) and v not in ("Thiếu/Không đọc được", "—"):
                    cell.number_format = '@'
                    cell.value = str(v)
                if c_idx == 5:
                    cell.fill = PatternFill("solid", fgColor=fill_c)
                    cell.font = bold

    _auto_fit_columns(ws6, len(h6))

    wb.save(out_path)
    return out_path
