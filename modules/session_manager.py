"""Quản lý nghiệp vụ Đợt thi, Lớp học, Học sinh, Mã đề và Bộ đáp án.

Đảm bảo:
- Mã đợt thi, mã đề, SBD là chuỗi giữ số 0 đầu ("01", "0023", "101").
- Cách ly hoàn toàn dữ liệu giữa các đợt thi.
- Hỗ trợ nhập danh sách học sinh & đáp án từ Excel / CSV.
- Tự động đối chiếu SBD, phát hiện trùng lặp, chuyển trạng thái NEEDS_REVIEW.
- Hỗ trợ chấm lại khi đáp án chuẩn thay đổi.
"""

import csv
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
from typing import Any

from .db import get_connection, init_db, transaction
from .errors import OMRError
from .scoring import ScoringError, score_answers, VALID_ANSWERS

try:
    import openpyxl
except ImportError:
    openpyxl = None

DEFAULT_GRADE_SCALE = {
    "gioi": 8.0,
    "kha": 6.5,
    "trung_binh": 5.0,
    "yeu": 3.5,
}


def classify_score(score: float, scale_config: dict[str, float] | None = None) -> str:
    """Xếp loại điểm số theo thang 10."""
    cfg = scale_config or DEFAULT_GRADE_SCALE
    if score >= cfg.get("gioi", 8.0):
        return "Giỏi"
    if score >= cfg.get("kha", 6.5):
        return "Khá"
    if score >= cfg.get("trung_binh", 5.0):
        return "Trung bình"
    if score >= cfg.get("yeu", 3.5):
        return "Yếu"
    return "Kém"


def compute_file_hash(path: str | Path) -> str | None:
    """Tính mã băm SHA-256 của file để nhận diện chính xác nội dung ảnh."""
    p = Path(path)
    if not p.is_file():
        return None
    try:
        h = hashlib.sha256()
        with open(p, "rb") as f:
            while chunk := f.read(65536):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return None


class SessionManager:
    """Lớp quản lý nghiệp vụ đợt thi, học sinh, mã đề và kết quả."""

    def __init__(self, db_path: Path | str | None = None):
        self.db_path = db_path
        init_db(self.db_path)

    # ── ĐỢT THI (EXAM SESSION) ────────────────────────────────────────────────

    def create_session(
        self,
        session_code: str,
        name: str,
        subject: str,
        exam_date: str,
        question_count: int = 50,
        max_score: float = 10.0,
        status: str = "ACTIVE",
        grade_scale_config: dict | None = None,
        class_codes: list[str] | None = None,
    ) -> dict:
        """Tạo đợt thi mới; session_code phải là chuỗi (bảo toàn số 0 đầu)."""
        session_code = str(session_code).strip()
        if not session_code:
            raise OMRError("Mã đợt thi không được để trống.")
        name = str(name).strip()
        if not name:
            raise OMRError("Tên đợt thi không được để trống.")
        if question_count != 50:
            raise OMRError(f"Mẫu phiếu thi hiện tại chỉ hỗ trợ đúng 50 câu (nhận được: {question_count}).")
        if max_score != 10.0:
            raise OMRError(f"Hệ thống hiện tại chỉ hỗ trợ thang điểm 10.0 (nhận được: {max_score}).")

        now = datetime.now().isoformat()
        scale_json = json.dumps(grade_scale_config or DEFAULT_GRADE_SCALE)

        with transaction(self.db_path) as conn:
            # Kiểm tra trùng mã
            cur = conn.execute("SELECT id FROM exam_sessions WHERE session_code = ?", (session_code,))
            if cur.fetchone():
                raise OMRError(f"Mã đợt thi '{session_code}' đã tồn tại.")

            conn.execute(
                """
                INSERT INTO exam_sessions (session_code, name, subject, exam_date, question_count,
                                           max_score, status, grade_scale_config, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (session_code, name, subject, exam_date, question_count, max_score, status, scale_json, now, now),
            )
            if class_codes:
                for cc in class_codes:
                    cc = str(cc).strip()
                    if cc:
                        # Đảm bảo lớp tồn tại
                        conn.execute("INSERT OR IGNORE INTO classes (class_code, class_name) VALUES (?, ?)", (cc, f"Lớp {cc}"))
                        conn.execute(
                            "INSERT OR IGNORE INTO session_classes (session_code, class_code) VALUES (?, ?)",
                            (session_code, cc),
                        )

        session = self.get_session(session_code)
        if session is None:
            raise OMRError(f"Không thể tải đợt thi vừa tạo '{session_code}'.")
        return session

    def update_session(self, session_code: str, **kwargs) -> dict:
        """Cập nhật thông tin đợt thi."""
        session_code = self.resolve_session_code(session_code)
        session = self.get_session(session_code)
        if session is None:
            raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")

        if "question_count" in kwargs and kwargs["question_count"] != 50:
            raise OMRError(f"Mẫu phiếu thi chỉ hỗ trợ 50 câu (nhận được: {kwargs['question_count']}).")
        if "max_score" in kwargs and kwargs["max_score"] != 10.0:
            raise OMRError(f"Hệ thống chỉ hỗ trợ thang điểm 10.0 (nhận được: {kwargs['max_score']}).")

        if session["status"] in ("LOCKED", "CLOSED") and kwargs.get("status") not in ("ACTIVE", None):
            raise OMRError(f"Đợt thi '{session_code}' đang ở trạng thái {session['status']}. Cần mở khóa trước khi chỉnh sửa thông tin.")

        allowed = {"name", "subject", "exam_date", "question_count", "max_score", "status", "grade_scale_config"}
        updates = []
        params = []
        for key, val in kwargs.items():
            if key in allowed:
                if key == "grade_scale_config" and isinstance(val, dict):
                    val = json.dumps(val)
                updates.append(f"{key} = ?")
                params.append(val)
        if not updates and "class_codes" not in kwargs:
            return session

        now = datetime.now().isoformat()
        updates.append("updated_at = ?")
        params.append(now)
        params.append(session_code)

        with transaction(self.db_path) as conn:
            if updates:
                sql = f"UPDATE exam_sessions SET {', '.join(updates)} WHERE session_code = ?"
                conn.execute(sql, params)

            if "class_codes" in kwargs and kwargs["class_codes"] is not None:
                conn.execute("DELETE FROM session_classes WHERE session_code = ?", (session_code,))
                for cc in kwargs["class_codes"]:
                    cc = str(cc).strip()
                    if cc:
                        conn.execute("INSERT OR IGNORE INTO classes (class_code, class_name) VALUES (?, ?)", (cc, f"Lớp {cc}"))
                        conn.execute("INSERT INTO session_classes (session_code, class_code) VALUES (?, ?)", (session_code, cc))

        session = self.get_session(session_code)
        if session is None:
            raise OMRError(f"Không thể tải đợt thi vừa cập nhật '{session_code}'.")
        return session

    def lock_session(self, session_code: str, reason: str = "", operator: str = "admin") -> dict:
        """Khóa đợt thi để không nhận bài mới và không cho phép sửa kết quả."""
        session_code = self.resolve_session_code(session_code)
        sess = self.get_session(session_code)
        if not sess:
            raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")
        now = datetime.now().isoformat()
        with transaction(self.db_path) as conn:
            conn.execute(
                "UPDATE exam_sessions SET status = 'LOCKED', updated_at = ? WHERE session_code = ?",
                (now, session_code),
            )
            conn.execute(
                """
                INSERT INTO session_audit_logs (session_code, action, reason, operator, details_json, created_at)
                VALUES (?, 'LOCK', ?, ?, ?, ?)
                """,
                (session_code, reason or "Khóa đợt thi", operator, json.dumps({"previous_status": sess["status"]}), now),
            )
        return self.get_session(session_code)

    def unlock_session(self, session_code: str, reason: str = "", operator: str = "admin") -> dict:
        """Mở khóa đợt thi (ghi nhận lý do và audit log)."""
        session_code = self.resolve_session_code(session_code)
        sess = self.get_session(session_code)
        if not sess:
            raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")
        now = datetime.now().isoformat()
        with transaction(self.db_path) as conn:
            conn.execute(
                "UPDATE exam_sessions SET status = 'ACTIVE', updated_at = ? WHERE session_code = ?",
                (now, session_code),
            )
            conn.execute(
                """
                INSERT INTO session_audit_logs (session_code, action, reason, operator, details_json, created_at)
                VALUES (?, 'UNLOCK', ?, ?, ?, ?)
                """,
                (session_code, reason or "Mở khóa đợt thi", operator, json.dumps({"previous_status": sess["status"]}), now),
            )
        return self.get_session(session_code)

    def close_session(self, session_code: str, reason: str = "", operator: str = "admin") -> dict:
        """Đóng đợt thi chính thức."""
        session_code = self.resolve_session_code(session_code)
        sess = self.get_session(session_code)
        if not sess:
            raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")
        now = datetime.now().isoformat()
        with transaction(self.db_path) as conn:
            conn.execute(
                "UPDATE exam_sessions SET status = 'CLOSED', updated_at = ? WHERE session_code = ?",
                (now, session_code),
            )
            conn.execute(
                """
                INSERT INTO session_audit_logs (session_code, action, reason, operator, details_json, created_at)
                VALUES (?, 'CLOSE', ?, ?, ?, ?)
                """,
                (session_code, reason or "Đóng đợt thi", operator, json.dumps({"previous_status": sess["status"]}), now),
            )
        return self.get_session(session_code)

    def get_session_audit_logs(self, session_code: str) -> list[dict]:
        """Lấy lịch sử thao tác khóa/mở khóa đợt thi."""
        session_code = self.resolve_session_code(session_code)
        conn = get_connection(self.db_path)
        cur = conn.execute(
            "SELECT * FROM session_audit_logs WHERE session_code = ? ORDER BY id ASC",
            (session_code,),
        )
        return [dict(r) for r in cur.fetchall()]

    def resolve_session_code(self, session_code: str | None) -> str:
        """Xác định session_code chuẩn xác nhất trong database.
        1. Ưu tiên khớp chính xác chuỗi session_code lưu trong DB (vd '01').
        2. Nếu không khớp và session_code có dạng số (vd '1'), tìm đợt thi tương ứng (vd '01' hoặc '001') trong DB.
        3. Trả về đúng chuỗi session_code đã lưu trong database nếu tìm thấy.
        4. Nếu không tìm thấy trong DB, trả về session_code gốc đã chuẩn hóa khoảng trắng.
        """
        if session_code is None:
            return ""
        s = str(session_code).strip()
        if not s:
            return ""
        conn = get_connection(self.db_path)
        cur = conn.execute("SELECT session_code FROM exam_sessions WHERE session_code = ?", (s,))
        row = cur.fetchone()
        if row:
            return str(row["session_code"])
        if s.isdigit():
            for alt in [s.zfill(2), s.zfill(3), s.lstrip("0")]:
                if alt and alt != s:
                    cur = conn.execute("SELECT session_code FROM exam_sessions WHERE session_code = ?", (alt,))
                    row = cur.fetchone()
                    if row:
                        return str(row["session_code"])
        return s

    def get_session(self, session_code: str) -> dict | None:
        """Lấy thông tin đợt thi theo mã (đồng bộ chuẩn hóa số 0 ở đầu)."""
        session_code = self.resolve_session_code(session_code)
        if not session_code:
            return None
        conn = get_connection(self.db_path)
        cur = conn.execute("SELECT * FROM exam_sessions WHERE session_code = ?", (session_code,))
        row = cur.fetchone()
        if not row:
            return None
        res = dict(row)
        if res.get("grade_scale_config"):
            try:
                res["grade_scale_config"] = json.loads(res["grade_scale_config"])
            except Exception:
                res["grade_scale_config"] = DEFAULT_GRADE_SCALE
        # Lấy danh sách lớp
        cur2 = conn.execute("SELECT class_code FROM session_classes WHERE session_code = ?", (session_code,))
        res["class_codes"] = [r["class_code"] for r in cur2.fetchall()]
        # Lấy danh sách mã đề
        cur3 = conn.execute("SELECT exam_code FROM exam_codes WHERE session_code = ?", (session_code,))
        res["exam_codes"] = [r["exam_code"] for r in cur3.fetchall()]
        return res

    def list_sessions(self, search: str | None = None, status: str | None = None) -> list[dict]:
        """Tìm kiếm và lọc danh sách đợt thi."""
        conn = get_connection(self.db_path)
        query = "SELECT * FROM exam_sessions WHERE 1=1"
        params: list[Any] = []
        if status:
            query += " AND status = ?"
            params.append(status)
        if search:
            query += " AND (session_code LIKE ? OR name LIKE ? OR subject LIKE ?)"
            s = f"%{search}%"
            params.extend([s, s, s])
        query += " ORDER BY session_code DESC"
        cur = conn.execute(query, params)
        rows = cur.fetchall()
        result = []
        for r in rows:
            d = dict(r)
            if d.get("grade_scale_config"):
                try:
                    d["grade_scale_config"] = json.loads(d["grade_scale_config"])
                except Exception:
                    d["grade_scale_config"] = DEFAULT_GRADE_SCALE
            result.append(d)
        return result

    # ── LỚP HỌC VÀ HỌC SINH ──────────────────────────────────────────────────

    def add_class(self, class_code: str, class_name: str | None = None) -> dict:
        """Thêm lớp học."""
        class_code = str(class_code).strip()
        class_name = str(class_name or f"Lớp {class_code}").strip()
        with transaction(self.db_path) as conn:
            conn.execute(
                "INSERT INTO classes (class_code, class_name) VALUES (?, ?) ON CONFLICT(class_code) DO UPDATE SET class_name=excluded.class_name",
                (class_code, class_name),
            )
        return {"class_code": class_code, "class_name": class_name}

    def list_classes(self) -> list[dict]:
        """Liệt kê danh sách lớp."""
        conn = get_connection(self.db_path)
        cur = conn.execute("SELECT * FROM classes ORDER BY class_code ASC")
        return [dict(r) for r in cur.fetchall()]

    def add_student(self, student_code: str, sbd: str, full_name: str, class_code: str) -> dict:
        """Thêm hoặc cập nhật học sinh."""
        student_code = str(student_code).strip()
        sbd = str(sbd).strip()
        full_name = str(full_name).strip()
        class_code = str(class_code).strip()
        if not student_code or not sbd or not full_name or not class_code:
            raise OMRError("Thông tin học sinh không được để trống trường nào.")

        with transaction(self.db_path) as conn:
            conn.execute("INSERT OR IGNORE INTO classes (class_code, class_name) VALUES (?, ?)", (class_code, f"Lớp {class_code}"))
            conn.execute(
                """
                INSERT INTO students (student_code, sbd, full_name, class_code)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(student_code) DO UPDATE SET
                    sbd=excluded.sbd,
                    full_name=excluded.full_name,
                    class_code=excluded.class_code
                """,
                (student_code, sbd, full_name, class_code),
            )
        return {"student_code": student_code, "sbd": sbd, "full_name": full_name, "class_code": class_code}

    def get_student_by_sbd(self, sbd: str, session_code: str | None = None) -> dict | None:
        """Tìm học sinh theo số báo danh.
        Nếu có session_code: CHỈ tìm trong các lớp thuộc đợt thi (session_classes).
        KHÔNG fallback sang các lớp ngoài đợt thi.
        Nếu phát hiện nhiều học sinh trùng SBD trong đợt thi, đánh dấu is_ambiguous = True.
        """
        sbd = str(sbd).strip()
        conn = get_connection(self.db_path)
        candidates = [sbd]
        if sbd.isdigit():
            z = sbd.zfill(6)
            if z not in candidates:
                candidates.append(z)
            ns = sbd.lstrip("0")
            if ns and ns not in candidates:
                candidates.append(ns)

        placeholders = ",".join("?" for _ in candidates)

        if session_code:
            cur_count = conn.execute("SELECT COUNT(*) as c FROM session_classes WHERE session_code = ?", (session_code,))
            has_class_restrictions = (cur_count.fetchone()["c"] > 0)
            if has_class_restrictions:
                cur = conn.execute(
                    f"""
                    SELECT s.* FROM students s
                    INNER JOIN session_classes sc ON s.class_code = sc.class_code
                    WHERE s.sbd IN ({placeholders}) AND sc.session_code = ?
                    """,
                    (*candidates, session_code),
                )
                rows = [dict(r) for r in cur.fetchall()]
                if not rows:
                    return None
                distinct_codes = {r["student_code"] for r in rows}
                res = dict(rows[0])
                if len(distinct_codes) > 1:
                    res["is_ambiguous"] = True
                return res

        cur = conn.execute(
            f"SELECT * FROM students WHERE sbd IN ({placeholders})",
            candidates,
        )
        rows = [dict(r) for r in cur.fetchall()]
        if not rows:
            return None
        distinct_codes = {r["student_code"] for r in rows}
        res = dict(rows[0])
        if len(distinct_codes) > 1:
            res["is_ambiguous"] = True
        return res

    def list_students(self, class_code: str | None = None, search: str | None = None) -> list[dict]:
        """Liệt kê danh sách học sinh theo lớp hoặc tìm kiếm."""
        conn = get_connection(self.db_path)
        query = "SELECT * FROM students WHERE 1=1"
        params: list[Any] = []
        if class_code:
            query += " AND class_code = ?"
            params.append(class_code)
        if search:
            query += " AND (student_code LIKE ? OR sbd LIKE ? OR full_name LIKE ?)"
            s = f"%{search}%"
            params.extend([s, s, s])
        query += " ORDER BY class_code ASC, sbd ASC"
        cur = conn.execute(query, params)
        return [dict(r) for r in cur.fetchall()]

    def import_students_from_csv(self, file_path: str | Path, default_class: str | None = None) -> int:
        """Nhập danh sách học sinh từ file CSV.

        Header hỗ trợ: student_code, sbd, full_name, class_code
        (hoặc tiếng Việt: ma_hs, sbd, ho_ten, lop).
        """
        path = Path(file_path)
        if not path.is_file():
            raise OMRError(f"File CSV không tồn tại: {path}")

        count = 0
        with open(path, mode="r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                # Chuẩn hóa tên cột
                r = {k.strip().lower(): str(v).strip() for k, v in row.items() if k}
                code = r.get("student_code") or r.get("ma_hs") or r.get("mã hs") or r.get("mahs")
                sbd = r.get("sbd") or r.get("so_bao_danh") or r.get("số báo danh")
                name = r.get("full_name") or r.get("ho_ten") or r.get("họ và tên") or r.get("hoten")
                c_code = r.get("class_code") or r.get("lop") or r.get("lớp") or default_class

                if code and sbd and name and c_code:
                    self.add_student(code, sbd, name, c_code)
                    count += 1
        return count

    def import_students_from_excel(self, file_path: str | Path, default_class: str | None = None) -> int:
        """Nhập danh sách học sinh từ file Excel .xlsx."""
        if openpyxl is None:
            raise OMRError("Cần thư viện openpyxl để nhập file Excel.")
        path = Path(file_path)
        if not path.is_file():
            raise OMRError(f"File Excel không tồn tại: {path}")

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        if ws is None:
            raise OMRError(f"File Excel không có trang tính đang hoạt động: {path}")
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return 0

        header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        # Tìm index các cột
        col_map = {}
        for idx, col in enumerate(header):
            if col in ("student_code", "ma_hs", "mã hs", "mahs", "mã học sinh"):
                col_map["code"] = idx
            elif col in ("sbd", "so_bao_danh", "số báo danh"):
                col_map["sbd"] = idx
            elif col in ("full_name", "ho_ten", "họ và tên", "hoten", "họ tên"):
                col_map["name"] = idx
            elif col in ("class_code", "lop", "lớp", "mã lớp"):
                col_map["class"] = idx

        count = 0
        for r in rows[1:]:
            if not any(r):
                continue
            code = str(r[col_map["code"]]).strip() if "code" in col_map and r[col_map["code"]] is not None else None
            sbd = str(r[col_map["sbd"]]).strip() if "sbd" in col_map and r[col_map["sbd"]] is not None else None
            name = str(r[col_map["name"]]).strip() if "name" in col_map and r[col_map["name"]] is not None else None
            c_code = str(r[col_map["class"]]).strip() if "class" in col_map and r[col_map["class"]] is not None else default_class

            # Chuyển đổi định dạng số sang chuỗi nguyên vẹn (ví dụ 101.0 -> 101)
            if code and code.endswith(".0"):
                code = code[:-2]
            if sbd and sbd.endswith(".0"):
                sbd = sbd[:-2]

            if code and sbd and name and c_code:
                self.add_student(code, sbd, name, c_code)
                count += 1
        return count

    # ── MÃ ĐỀ VÀ BỘ ĐÁP ÁN (EXAM CODES & ANSWER KEYS) ─────────────────────────

    def set_answer_key(
        self,
        session_code: str,
        exam_code: str,
        answers: list[str],
        source: str = "manual",
        overwrite: bool = True,
    ) -> dict | None:
        """Lưu hoặc cập nhật bộ đáp án cho mã đề trong đợt thi."""
        session_code = self.resolve_session_code(session_code)
        exam_code = str(exam_code).strip()
        if not session_code or not exam_code:
            raise OMRError("Mã đợt thi và mã đề không được để trống.")

        # Lấy session để biết question_count
        sess = self.get_session(session_code)
        if not sess:
            raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")
        if sess.get("status") in ("LOCKED", "CLOSED"):
            raise OMRError(f"Đợt thi '{session_code}' đang ở trạng thái {sess.get('status')}. Không thể thay đổi đáp án chuẩn.")
        expected_len = sess["question_count"]

        # Nếu không cho phép ghi đè và đã có đáp án, bỏ qua
        if not overwrite and self.get_answer_key(session_code, exam_code) is not None:
            return None

        if len(answers) != expected_len:
            raise ScoringError(f"Đáp án mã đề {exam_code} cần đúng {expected_len} câu, nhận được {len(answers)}.")

        for i, ans in enumerate(answers, start=1):
            if ans not in VALID_ANSWERS:
                raise ScoringError(f"Câu {i}: đáp án '{ans}' không hợp lệ (chỉ nhận A, B, C, D).")

        now = datetime.now().isoformat()
        ans_json = json.dumps(answers)

        with transaction(self.db_path) as conn:
            conn.execute(
                """
                INSERT INTO exam_codes (session_code, exam_code, answers_json, source, created_at)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(session_code, exam_code) DO UPDATE SET
                    answers_json=excluded.answers_json,
                    source=excluded.source
                """,
                (session_code, exam_code, ans_json, source, now),
            )

        return {"session_code": session_code, "exam_code": exam_code, "answers": answers, "source": source}

    def get_answer_key(self, session_code: str, exam_code: str) -> list[str] | None:
        """Lấy bộ đáp án đúng theo Mã đợt thi + Mã đề."""
        session_code = self.resolve_session_code(session_code)
        exam_code = str(exam_code).strip()
        conn = get_connection(self.db_path)
        cur = conn.execute(
            "SELECT answers_json FROM exam_codes WHERE session_code = ? AND exam_code = ?",
            (session_code, exam_code),
        )
        row = cur.fetchone()
        if not row:
            return None
        return json.loads(row["answers_json"])

    def list_exam_codes(self, session_code: str) -> list[dict]:
        """Liệt kê tất cả các mã đề trong một đợt thi."""
        session_code = self.resolve_session_code(session_code)
        conn = get_connection(self.db_path)
        cur = conn.execute(
            "SELECT session_code, exam_code, source, created_at, answers_json FROM exam_codes WHERE session_code = ? ORDER BY exam_code ASC",
            (session_code,),
        )
        rows = cur.fetchall()
        result = []
        for r in rows:
            d = dict(r)
            d["answers"] = json.loads(d.pop("answers_json"))
            result.append(d)
        return result

    def delete_exam_code(self, session_code: str, exam_code: str) -> bool:
        """Xóa một mã đề khỏi đợt thi qua transaction an toàn."""
        session_code = self.resolve_session_code(session_code)
        exam_code = str(exam_code).strip()
        sess = self.get_session(session_code)
        if not sess:
            raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")
        if sess.get("status") in ("LOCKED", "CLOSED"):
            raise OMRError(f"Đợt thi '{session_code}' đang ở trạng thái {sess.get('status')}. Không thể xóa mã đề.")
        with transaction(self.db_path) as conn:
            cur = conn.execute(
                "DELETE FROM exam_codes WHERE session_code = ? AND exam_code = ?",
                (session_code, exam_code),
            )
            return cur.rowcount > 0

    def import_answer_keys_from_csv(self, session_code: str | None, file_path: str | Path) -> int:
        """Nhập nhiều mã đề và đáp án từ file CSV.

        Hỗ trợ các định dạng:
        - 4 cột: session_code, exam_code, question, answer
        - 3 cột: exam_code, question, answer
        - 2 cột: question, answer (tự suy ra mã đề từ tên file hoặc tham số)
        - Dòng ngang: exam_code, 1, 2, ..., 50
        """
        path = Path(file_path)
        if not path.is_file():
            raise OMRError(f"File CSV không tồn tại: {path}")

        # Tự suy đoán mã đề hoặc đợt thi từ tên file nếu chưa có
        inferred_session = session_code
        inferred_exam = None
        match_sess = re.search(r"dot_?(\d+)", path.stem, re.IGNORECASE)
        if match_sess and not inferred_session:
            inferred_session = match_sess.group(1).zfill(2)
        match_code = re.search(r"ma_?(\d+)", path.stem, re.IGNORECASE)
        if match_code:
            inferred_exam = match_code.group(1)

        count = 0
        with open(path, mode="r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            header = [c.strip().lower() for c in next(reader, [])]
            if not header:
                return 0

            # Xác định vị trí các cột nếu có header
            col_sess_idx = -1
            col_exam_idx = -1
            col_q_idx = -1
            col_ans_idx = -1
            for idx, col in enumerate(header):
                if col in ("session_code", "dot_thi", "ma_dot_thi", "đợt thi", "mã đợt thi"):
                    col_sess_idx = idx
                elif col in ("exam_code", "ma_de", "mã đề", "made"):
                    col_exam_idx = idx
                elif col in ("question", "cau", "so_cau", "câu", "số thứ tự", "stt"):
                    col_q_idx = idx
                elif col in ("answer", "dap_an", "đáp án", "dapan"):
                    col_ans_idx = idx

            rows = list(reader)

            # Trường hợp dòng ngang: cột 1 là exam_code, các cột tiếp theo là đáp án 1..50
            if len(header) >= 51 and col_q_idx == -1:
                sess_target = inferred_session or session_code
                if not sess_target:
                    raise OMRError("Cần cung cấp mã đợt thi để nhập đáp án.")
                sess = self.get_session(sess_target)
                if not sess:
                    raise OMRError(f"Đợt thi '{sess_target}' không tồn tại.")
                q_count = sess["question_count"]
                for row in rows:
                    if not row or not any(row):
                        continue
                    ecode = str(row[0]).strip()
                    if ecode.endswith(".0"):
                        ecode = ecode[:-2]
                    ans = [str(x).strip().upper() for x in row[1 : q_count + 1] if str(x).strip()]
                    if len(ans) == q_count and all(a in VALID_ANSWERS for a in ans):
                        self.set_answer_key(sess_target, ecode, ans, source="csv")
                        count += 1
                return count

            # Trường hợp bảng dọc: nhóm theo (session_code, exam_code) -> {question: answer}
            grouped: dict[tuple[str, str], dict[int, str]] = {}
            for r in rows:
                if not r or not any(r):
                    continue
                # Trích xuất session, exam, q, ans
                s_val = None
                e_val = None
                q_val = None
                a_val = None

                if col_sess_idx != -1 and len(r) > col_sess_idx:
                    s_val = str(r[col_sess_idx]).strip()
                if col_exam_idx != -1 and len(r) > col_exam_idx:
                    e_val = str(r[col_exam_idx]).strip()
                if col_q_idx != -1 and len(r) > col_q_idx:
                    q_val = str(r[col_q_idx]).strip()
                if col_ans_idx != -1 and len(r) > col_ans_idx:
                    a_val = str(r[col_ans_idx]).strip().upper()

                # Nếu không tìm thấy bằng header, dựa vào số cột
                if col_ans_idx == -1:
                    if len(r) >= 4:
                        s_val = str(r[0]).strip()
                        e_val = str(r[1]).strip()
                        q_val = str(r[2]).strip()
                        a_val = str(r[3]).strip().upper()
                    elif len(r) == 3:
                        e_val = str(r[0]).strip()
                        q_val = str(r[1]).strip()
                        a_val = str(r[2]).strip().upper()
                    elif len(r) == 2:
                        q_val = str(r[0]).strip()
                        a_val = str(r[1]).strip().upper()

                s_val = s_val or inferred_session or session_code
                e_val = e_val or inferred_exam
                if e_val and e_val.endswith(".0"):
                    e_val = e_val[:-2]

                if not s_val or not e_val or not q_val or not a_val:
                    continue

                try:
                    q_int = int(q_val)
                    if a_val in VALID_ANSWERS:
                        grouped.setdefault((s_val, e_val), {})[q_int] = a_val
                except ValueError:
                    continue

            for (s_code, e_code), ans_map in grouped.items():
                sess = self.get_session(s_code)
                if not sess:
                    continue
                q_count = sess["question_count"]
                if len(ans_map) == q_count and all(i in ans_map for i in range(1, q_count + 1)):
                    ans_list = [ans_map[i] for i in range(1, q_count + 1)]
                    self.set_answer_key(s_code, e_code, ans_list, source="csv")
                    count += 1

        return count

    def import_answer_keys_from_json(
        self,
        file_path: str | Path,
        default_session_code: str | None = None,
        overwrite_existing: bool = True,
    ) -> int:
        """Nhập bộ đáp án từ file JSON (hỗ trợ JSON tổng hợp đa đợt thi hoặc đơn lẻ)."""
        path = Path(file_path)
        if not path.is_file():
            raise OMRError(f"File JSON không tồn tại: {path}")
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            raise OMRError(f"Không thể đọc file JSON: {exc}") from exc

        count = 0
        # Dạng 1: {"01": {"101": ["A", ...], "102": [...]}, "02": {...}}
        if isinstance(data, dict):
            # Kiểm tra xem có phải dạng nested session -> exam -> answers
            first_val = next(iter(data.values())) if data else None
            if isinstance(first_val, dict):
                for s_code, exams in data.items():
                    s_code = str(s_code).strip()
                    sess = self.get_session(s_code)
                    if not sess:
                        continue
                    if isinstance(exams, dict):
                        for e_code, ans_list in exams.items():
                            if isinstance(ans_list, list) and len(ans_list) == sess["question_count"]:
                                res = self.set_answer_key(s_code, str(e_code).strip(), ans_list, source="json", overwrite=overwrite_existing)
                                if res is not None:
                                    count += 1
                return count

            # Dạng 2: {"101": ["A", ...], "102": [...]} với default_session_code
            if isinstance(first_val, list):
                s_code = default_session_code
                if not s_code:
                    raise OMRError("Cần chỉ định mã đợt thi khi nhập file JSON chứa danh sách mã đề.")
                sess = self.get_session(s_code)
                if not sess:
                    raise OMRError(f"Đợt thi '{s_code}' không tồn tại.")
                for e_code, ans_list in data.items():
                    if isinstance(ans_list, list) and len(ans_list) == sess["question_count"]:
                        res = self.set_answer_key(s_code, str(e_code).strip(), ans_list, source="json", overwrite=overwrite_existing)
                        if res is not None:
                            count += 1
                return count

            # Dạng 3: {"session_code": "01", "exam_code": "101", "answers": [...]}
            if "answers" in data and ("exam_code" in data or "session_code" in data):
                s_code = str(data.get("session_code") or default_session_code or "").strip()
                e_code = str(data.get("exam_code") or "").strip()
                ans_list = data["answers"]
                if s_code and e_code and isinstance(ans_list, list):
                    res = self.set_answer_key(s_code, e_code, ans_list, source="json", overwrite=overwrite_existing)
                    return 1 if res is not None else 0

        # Dạng 4: List of dicts [{"session_code": "01", "exam_code": "101", "answers": [...]}, ...]
        elif isinstance(data, list):
            for item in data:
                if isinstance(item, dict) and "answers" in item:
                    s_code = str(item.get("session_code") or default_session_code or "").strip()
                    e_code = str(item.get("exam_code") or "").strip()
                    ans_list = item["answers"]
                    if s_code and e_code and isinstance(ans_list, list):
                        sess = self.get_session(s_code)
                        if sess and len(ans_list) == sess["question_count"]:
                            res = self.set_answer_key(s_code, e_code, ans_list, source="json", overwrite=overwrite_existing)
                            if res is not None:
                                count += 1

        return count

    def import_answer_keys_from_excel(self, session_code: str, file_path: str | Path) -> int:
        """Nhập nhiều mã đề và đáp án từ file Excel .xlsx."""
        if openpyxl is None:
            raise OMRError("Cần openpyxl để nhập file Excel.")
        path = Path(file_path)
        if not path.is_file():
            raise OMRError(f"File Excel không tồn tại: {path}")

        sess = self.get_session(session_code)
        if not sess:
            raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")
        q_count = sess["question_count"]

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        if ws is None:
            raise OMRError(f"File Excel không có trang tính đang hoạt động: {path}")
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return 0

        count = 0
        header = rows[0]
        # Kiểm tra xem dạng ngang hay dọc
        if len(header) >= q_count + 1:
            for r in rows[1:]:
                if not any(r):
                    continue
                ecode = str(r[0]).strip()
                if ecode.endswith(".0"):
                    ecode = ecode[:-2]
                ans = [str(x).strip().upper() for x in r[1 : q_count + 1] if x is not None]
                if len(ans) == q_count and all(a in VALID_ANSWERS for a in ans):
                    self.set_answer_key(session_code, ecode, ans, source="excel")
                    count += 1
        return count

    # ── LƯU VÀ CHẤM LẠI KẾT QUẢ (RESULTS & RE-GRADING) ─────────────────────────

    def save_exam_result(
        self,
        session_code: str,
        image_path: str | Path,
        student_answers: list[str],
        sbd: str | None = None,
        exam_code: str | None = None,
        student_code: str | None = None,
        student_name: str | None = None,
        class_code: str | None = None,
        extra_review_reasons: list[str] | None = None,
    ) -> dict:
        """Lưu kết quả bài thi. Tự động kiểm tra trùng lặp và chấm điểm."""
        session_code = self.resolve_session_code(session_code)
        sess = self.get_session(session_code)
        if not sess:
            raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")

        if sess.get("status") in ("LOCKED", "CLOSED"):
            raise OMRError(f"Đợt thi '{session_code}' đang ở trạng thái {sess.get('status')}. Không thể nhận bài mới hoặc chỉnh sửa kết quả.")

        sbd = str(sbd).strip() if sbd else None
        exam_code = str(exam_code).strip() if exam_code else None
        canonical_path = Path(image_path).resolve().as_posix()
        file_hash = compute_file_hash(image_path)
        now = datetime.now().isoformat()

        # Đối chiếu học sinh nếu chưa có thông tin
        is_ambiguous = False
        if sbd and (not student_name or not class_code):
            stu = self.get_student_by_sbd(sbd, session_code)
            if stu:
                student_code = student_code or stu["student_code"]
                student_name = student_name or stu["full_name"]
                class_code = class_code or stu["class_code"]
                if stu.get("is_ambiguous"):
                    is_ambiguous = True

        # Kiểm tra trạng thái và các điều kiện
        review_reasons = []
        status = "SUCCESS"

        if extra_review_reasons:
            for err in extra_review_reasons:
                if err and err not in review_reasons:
                    review_reasons.append(err)
            status = "NEEDS_REVIEW"

        if is_ambiguous:
            status = "NEEDS_REVIEW"
            review_reasons.append(f"Phát hiện nhiều học sinh có cùng SBD '{sbd}' trong đợt thi.")

        if not exam_code:
            status = "NEEDS_REVIEW"
            review_reasons.append("Thiếu mã đề thi.")

        if not sbd:
            status = "NEEDS_REVIEW"
            review_reasons.append("Thiếu số báo danh.")
        elif not student_name:
            status = "NEEDS_REVIEW"
            review_reasons.append(f"Không tìm thấy học sinh với SBD '{sbd}' trong danh sách.")

        # Kiểm tra bài đã tồn tại trong đợt thi:
        # CHỈ khớp khi cùng canonical_path hoặc cùng image_hash.
        # KHÔNG sử dụng image_path LIKE %fname để tránh 2 ảnh khác thư mục trùng tên bị ghi đè.
        conn = get_connection(self.db_path)
        if file_hash:
            cur = conn.execute(
                """
                SELECT id, image_path FROM exam_results
                WHERE session_code = ? AND (image_path = ? OR image_hash = ?)
                ORDER BY id DESC LIMIT 1
                """,
                (session_code, canonical_path, file_hash),
            )
        else:
            cur = conn.execute(
                """
                SELECT id, image_path FROM exam_results
                WHERE session_code = ? AND image_path = ?
                ORDER BY id DESC LIMIT 1
                """,
                (session_code, canonical_path),
            )
        existing_row = cur.fetchone()
        existing_id = existing_row["id"] if existing_row else None

        # Kiểm tra trùng SBD trong đợt thi (loại trừ chính bài đang cập nhật lại cùng ảnh/hash)
        if sbd:
            cur = conn.execute(
                """
                SELECT id, image_path FROM exam_results
                WHERE session_code = ? AND sbd = ? AND status != 'FAILED'
                AND id != ?
                """,
                (session_code, sbd, existing_id or -1),
            )
            existing = cur.fetchall()
            if existing:
                status = "NEEDS_REVIEW"
                review_reasons.append(f"Trùng số báo danh '{sbd}' trong đợt thi (đã có {len(existing)} bài khác).")

        # Lấy đáp án chuẩn và chấm điểm
        score_res = None
        classification = None
        if exam_code:
            key = self.get_answer_key(session_code, exam_code)
            if not key:
                status = "NEEDS_REVIEW"
                review_reasons.append(f"Chưa có đáp án chuẩn cho mã đề '{exam_code}'.")
            else:
                score_res = score_answers(student_answers, key)
                classification = classify_score(score_res["score"], sess.get("grade_scale_config"))

        review_reason_str = "; ".join(review_reasons) if review_reasons else None

        with transaction(self.db_path) as conn:
            if existing_id is not None:
                conn.execute(
                    """
                    UPDATE exam_results SET
                        image_path = ?, image_hash = ?, student_code = ?, sbd = ?,
                        student_name = ?, class_code = ?, exam_code = ?, student_answers_json = ?,
                        score = ?, correct_count = ?, wrong_count = ?, blank_count = ?,
                        multiple_count = ?, total_questions = ?, details_json = ?,
                        classification = ?, status = ?, review_reason = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        canonical_path,
                        file_hash,
                        student_code,
                        sbd,
                        student_name,
                        class_code,
                        exam_code,
                        json.dumps(student_answers),
                        score_res["score"] if score_res else None,
                        score_res["correct_count"] if score_res else None,
                        score_res["wrong_count"] if score_res else None,
                        score_res["blank_count"] if score_res else None,
                        score_res["multiple_count"] if score_res else None,
                        sess["question_count"],
                        json.dumps(score_res["details"]) if score_res else None,
                        classification,
                        status,
                        review_reason_str,
                        now,
                        existing_id,
                    ),
                )
                result_id = existing_id
            else:
                cur = conn.execute(
                    """
                    INSERT INTO exam_results (
                        session_code, image_path, image_hash, student_code, sbd, student_name, class_code,
                        exam_code, student_answers_json, score, correct_count, wrong_count,
                        blank_count, multiple_count, total_questions, details_json,
                        classification, status, review_reason, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        session_code,
                        canonical_path,
                        file_hash,
                        student_code,
                        sbd,
                        student_name,
                        class_code,
                        exam_code,
                        json.dumps(student_answers),
                        score_res["score"] if score_res else None,
                        score_res["correct_count"] if score_res else None,
                        score_res["wrong_count"] if score_res else None,
                        score_res["blank_count"] if score_res else None,
                        score_res["multiple_count"] if score_res else None,
                        sess["question_count"],
                        json.dumps(score_res["details"]) if score_res else None,
                        classification,
                        status,
                        review_reason_str,
                        now,
                        now,
                    ),
                )
                result_id = cur.lastrowid

        if result_id is None:
            raise OMRError("Không thể xác định ID của kết quả bài thi vừa lưu.")
        result = self.get_result_by_id(result_id)
        if result is None:
            raise OMRError(f"Không thể tải kết quả bài thi vừa lưu (ID {result_id}).")
        return result

    def get_result_by_id(self, result_id: int) -> dict | None:
        """Lấy một kết quả bài thi theo ID."""
        conn = get_connection(self.db_path)
        cur = conn.execute("SELECT * FROM exam_results WHERE id = ?", (result_id,))
        row = cur.fetchone()
        if not row:
            return None
        res = dict(row)
        res["student_answers"] = json.loads(res["student_answers_json"])
        if res.get("details_json"):
            res["details"] = json.loads(res["details_json"])
        return res

    def list_results(
        self,
        session_code: str,
        class_code: str | None = None,
        exam_code: str | None = None,
        status: str | None = None,
    ) -> list[dict]:
        """Lấy danh sách kết quả bài thi theo bộ lọc."""
        session_code = self.resolve_session_code(session_code)
        conn = get_connection(self.db_path)
        query = "SELECT * FROM exam_results WHERE session_code = ?"
        params: list[Any] = [session_code]
        if class_code:
            query += " AND class_code = ?"
            params.append(class_code)
        if exam_code:
            query += " AND exam_code = ?"
            params.append(exam_code)
        if status:
            query += " AND status = ?"
            params.append(status)
        query += " ORDER BY id ASC"

        cur = conn.execute(query, params)
        rows = cur.fetchall()
        result = []
        for r in rows:
            d = dict(r)
            d["student_answers"] = json.loads(d["student_answers_json"])
            if d.get("details_json"):
                d["details"] = json.loads(d["details_json"])
            result.append(d)
        return result

    def re_grade_session(self, session_code: str) -> int:
        """Chấm lại toàn bộ bài thi trong đợt thi khi đáp án chuẩn hoặc ngưỡng điểm thay đổi."""
        session_code = self.resolve_session_code(session_code)
        sess = self.get_session(session_code)
        if not sess:
            raise OMRError(f"Đợt thi '{session_code}' không tồn tại.")

        if sess.get("status") in ("LOCKED", "CLOSED"):
            raise OMRError(f"Đợt thi '{session_code}' đang ở trạng thái {sess.get('status')}. Không thể chấm lại kết quả.")

        results = self.list_results(session_code)
        now = datetime.now().isoformat()
        updated_count = 0

        # Đếm số lượng SBD để kiểm tra bài trùng SBD trong đợt
        sbd_counts: dict[str, int] = {}
        for r in results:
            sbd_val = r.get("sbd")
            if sbd_val and r.get("status") != "FAILED":
                sbd_counts[sbd_val] = sbd_counts.get(sbd_val, 0) + 1

        with transaction(self.db_path) as conn:
            for item in results:
                ecode = item.get("exam_code")
                if not ecode:
                    continue
                key = self.get_answer_key(session_code, ecode)

                # Bảo toàn các lý do review trước đó (bỏ qua thông báo thiếu đáp án chuẩn hoặc trùng SBD cũ để tính lại)
                existing_reasons = []
                if item.get("review_reason"):
                    for part in item["review_reason"].split(";"):
                        p = part.strip()
                        if (
                            p
                            and not p.startswith("Chưa có đáp án chuẩn")
                            and not p.startswith("Không tìm thấy đáp án chuẩn")
                            and not p.startswith("Trùng số báo danh")
                        ):
                            existing_reasons.append(p)

                review_reasons = list(existing_reasons)
                new_status = "SUCCESS"

                sbd_val = item.get("sbd")
                if not sbd_val:
                    new_status = "NEEDS_REVIEW"
                    review_reasons.append("Thiếu số báo danh.")
                elif not item.get("student_name"):
                    new_status = "NEEDS_REVIEW"
                    review_reasons.append(f"Không tìm thấy học sinh với SBD '{sbd_val}' trong danh sách.")
                elif sbd_counts.get(sbd_val, 0) > 1:
                    new_status = "NEEDS_REVIEW"
                    review_reasons.append(f"Trùng số báo danh '{sbd_val}' trong đợt thi (đã có {sbd_counts[sbd_val] - 1} bài khác).")

                if not key:
                    new_status = "NEEDS_REVIEW"
                    review_reasons.append(f"Chưa có đáp án chuẩn cho mã đề '{ecode}'.")
                    score_val = None
                    correct_val = None
                    wrong_val = None
                    blank_val = None
                    mult_val = None
                    details_val = None
                    classification = None
                else:
                    student_answers = item["student_answers"]
                    score_res = score_answers(student_answers, key)
                    classification = classify_score(score_res["score"], sess.get("grade_scale_config"))
                    score_val = score_res["score"]
                    correct_val = score_res["correct_count"]
                    wrong_val = score_res["wrong_count"]
                    blank_val = score_res["blank_count"]
                    mult_val = score_res["multiple_count"]
                    details_val = json.dumps(score_res["details"])

                if review_reasons:
                    new_status = "NEEDS_REVIEW"

                conn.execute(
                    """
                    UPDATE exam_results
                    SET score = ?, correct_count = ?, wrong_count = ?, blank_count = ?,
                        multiple_count = ?, details_json = ?, classification = ?,
                        status = ?, review_reason = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        score_val,
                        correct_val,
                        wrong_val,
                        blank_val,
                        mult_val,
                        details_val,
                        classification,
                        new_status,
                        "; ".join(review_reasons) if review_reasons else None,
                        now,
                        item["id"],
                    ),
                )
                updated_count += 1

        return updated_count
