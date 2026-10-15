"""Xuất kết quả chấm điểm ra file Excel — Phase 7."""

from pathlib import Path

from .errors import OMRError

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False


class ExportError(OMRError):
    """Lỗi khi xuất file Excel."""


# Màu sắc
_COLOR_CORRECT  = "C6EFCE"   # xanh lá nhạt
_COLOR_WRONG    = "FFC7CE"   # đỏ nhạt
_COLOR_BLANK    = "FFEB9C"   # vàng nhạt
_COLOR_MULTIPLE = "BDD7EE"   # xanh lam nhạt
_COLOR_HEADER   = "4472C4"   # xanh đậm
_COLOR_SUMMARY  = "D9E1F2"   # xanh nhạt tổng kết


def _thin_border() -> "Border":
    side = Side(style="thin", color="AAAAAA")
    return Border(left=side, right=side, top=side, bottom=side)


def export_excel(score_result: dict, output_path: str | Path) -> Path:
    """Xuất 50 câu và thống kê ra file .xlsx.

    Args:
        score_result: dict trả về từ ``score_answers()``.
        output_path: đường dẫn file .xlsx sẽ ghi (tạo thư mục cha nếu thiếu).

    Returns:
        Path thực tế đã ghi.

    Raises:
        ExportError: nếu thiếu openpyxl, dữ liệu không hợp lệ hoặc lỗi ghi file.
    """
    if not _HAS_OPENPYXL:
        raise ExportError(
            "Thư viện openpyxl chưa cài. Chạy: pip install openpyxl"
        )
    required = {"score", "correct_count", "wrong_count", "blank_count",
                "multiple_count", "total", "details"}
    if not required.issubset(score_result):
        raise ExportError("score_result thiếu trường; hãy dùng kết quả từ score_answers().")

    output_path = Path(output_path)
    if output_path.suffix.lower() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kết quả"

    # ── Tiêu đề cột ──────────────────────────────────────────
    headers = ["Câu", "Đáp án HS", "Đáp án đúng", "Kết quả"]
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor=_COLOR_HEADER)
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = _thin_border()

    ws.row_dimensions[1].height = 22

    # ── Dữ liệu 50 câu ───────────────────────────────────────
    result_label = {
        "correct":  "✓ Đúng",
        "wrong":    "✗ Sai",
        "blank":    "— Bỏ trống",
        "multiple": "⚠ Nhiều đáp án",
    }
    result_color = {
        "correct":  _COLOR_CORRECT,
        "wrong":    _COLOR_WRONG,
        "blank":    _COLOR_BLANK,
        "multiple": _COLOR_MULTIPLE,
    }
    center = Alignment(horizontal="center", vertical="center")

    for detail in score_result["details"]:
        row = detail["question"] + 1
        result = detail["result"]
        fill = PatternFill("solid", fgColor=result_color.get(result, "FFFFFF"))
        border = _thin_border()

        values = [
            detail["question"],
            detail["student"],
            detail["correct"],
            result_label.get(result, result),
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.alignment = center
            cell.border = border
            if col_idx >= 2:  # tô màu từ cột "Đáp án HS"
                cell.fill = fill

    # ── Tổng kết ─────────────────────────────────────────────
    summary_row = score_result["total"] + 3
    summary_fill = PatternFill("solid", fgColor=_COLOR_SUMMARY)
    bold = Font(bold=True, name="Calibri", size=11)

    summary_data = [
        ("Tổng số câu",      score_result["total"]),
        ("Số câu đúng",      score_result["correct_count"]),
        ("Số câu sai",       score_result["wrong_count"]),
        ("Số câu bỏ trống",  score_result["blank_count"]),
        ("Nhiều đáp án",     score_result["multiple_count"]),
        ("Điểm (thang 10)",  score_result["score"]),
    ]
    for offset, (label, value) in enumerate(summary_data):
        label_cell = ws.cell(row=summary_row + offset, column=1, value=label)
        value_cell = ws.cell(row=summary_row + offset, column=2, value=value)
        for cell in (label_cell, value_cell):
            cell.font = bold
            cell.fill = summary_fill
            cell.border = _thin_border()
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Điểm → làm nổi bật
    score_cell = ws.cell(row=summary_row + 5, column=2)
    score_cell.font = Font(bold=True, size=14, color="C00000", name="Calibri")

    # ── Độ rộng cột ──────────────────────────────────────────
    col_widths = [8, 14, 16, 18]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── Ghi file ─────────────────────────────────────────────
    try:
        wb.save(output_path)
    except PermissionError as exc:
        raise ExportError(
            f"Không ghi được file Excel (có thể đang mở trong Excel): {output_path}\n{exc}"
        ) from exc
    except OSError as exc:
        raise ExportError(f"Lỗi ghi file Excel {output_path}: {exc}") from exc

    return output_path
